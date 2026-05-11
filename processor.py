import os
import pandas as pd
import openpyxl
import win32com.client
from sqlalchemy import create_engine
import urllib


class FormProcessor:
    def __init__(self, logger):
        self.logger = logger

    def load_data_sql(self):
        """
        Conecta a SQL Server y ejecuta la consulta.
        Solo cambia las credenciales en db_config.
        """
        try:
            db_config = {
                "host": "localhost",      # Ej: localhost o localhost\\SQLEXPRESS
                "db": "restrepo",
                "user": "sa",
                "pass": "TU_PASSWORD"
            }

            query = """
            SELECT
                s.id_solicitud,
                c.numero                          AS num_cliente,
                c.nombres_apellidos,
                c.cedula,
                c.ciudad_ng                       AS ciudad_ng,
                fv.actividad,
                fv.nombre_micro,
                s.inicial_actual                  AS monto_aprobado,
                s.meses_actual                    AS plazo,
                s.cuota_actual                    AS cuota_actual,
                s.ventas_actual,
                s.liq_disp_actual,
                s.estabilidad,
                s.habilidad_empresarial,
                cr.dias_mora,
                cr.analista_actual,
                fv.observaciones                  AS obs_visita,
                s.observaciones                   AS obs_solicitud
            FROM solicitudes s
            JOIN clientes c ON c.id_cliente = s.id_cliente
            LEFT JOIN fichas_visita fv ON fv.id_solicitud = s.id_solicitud
            LEFT JOIN creditos cr ON cr.id_solicitud = s.id_solicitud
                                   AND cr.monto_actual = s.inicial_actual
            ORDER BY s.id_solicitud
            """

            host = db_config["host"]
            db = db_config["db"]
            user = db_config["user"]
            password = db_config["pass"]

            params = urllib.parse.quote_plus(
                f"DRIVER={{ODBC Driver 17 for SQL Server}};"
                f"SERVER={host};"
                f"DATABASE={db};"
                f"UID={user};"
                f"PWD={password};"
                f"TrustServerCertificate=yes;"
            )
            connection_string = f"mssql+pyodbc:///?odbc_connect={params}"

            self.logger.info("Conectando a SQL Server...")
            engine = create_engine(connection_string)

            self.logger.info("Ejecutando consulta SQL...")
            df = pd.read_sql_query(query, engine)

            df.columns = df.columns.str.strip()
            df = df.fillna("")

            return df

        except Exception as e:
            self.logger.error(f"Error al conectar o consultar BD: {str(e)}")
            raise e

    def process_records(self, df, template_path, output_dir, progress_callback=None, log_callback=None):
        """
        Itera sobre el dataframe, llena el template y exporta a PDF.
        """
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
            self.logger.info(f"Directorio de salida creado: {output_dir}")

        total = len(df)
        success_count = 0

        try:
            excel_app = win32com.client.DispatchEx("Excel.Application")
            excel_app.Visible = False
            excel_app.DisplayAlerts = False
        except Exception as e:
            self.logger.error("No se pudo iniciar Microsoft Excel. Asegúrese de que está instalado.")
            if log_callback:
                log_callback("Error crítico: Microsoft Excel no encontrado o no se puede ejecutar.\n")
            raise e

        try:
            for index, row in df.iterrows():
                try:
                    name_parts = []
                    for field in ['nombres_apellidos', 'cedula', 'id_solicitud']:
                        matching_cols = [c for c in df.columns if field.lower() in c.lower()]
                        if matching_cols:
                            val = str(row[matching_cols[0]]).strip()
                            if val:
                                name_parts.append(val)

                    if not name_parts:
                        base_name = f"usuario_{index+1}"
                    else:
                        base_name = "_".join(name_parts)

                    base_name = "".join(c for c in base_name if c.isalnum() or c in (' ', '_', '-')).replace(" ", "_")

                    pdf_filename = f"{base_name}.pdf"
                    pdf_path = os.path.abspath(os.path.join(output_dir, pdf_filename))

                    temp_excel_filename = f"temp_{base_name}.xlsx"
                    temp_excel_path = os.path.abspath(os.path.join(output_dir, temp_excel_filename))

                    msg = f"Procesando: {pdf_filename}..."
                    self.logger.info(msg)
                    if log_callback:
                        log_callback(msg + "\n")

                    self._fill_excel_template(template_path, temp_excel_path, row)
                    self._convert_to_pdf(excel_app, temp_excel_path, pdf_path)

                    if os.path.exists(temp_excel_path):
                        os.remove(temp_excel_path)

                    success_count += 1

                except Exception as e:
                    err_msg = f"Error al procesar la fila {index+1}: {str(e)}"
                    self.logger.error(err_msg)
                    if log_callback:
                        log_callback(err_msg + "\n")

                if progress_callback:
                    progress_callback(index + 1, total)

        finally:
            excel_app.Quit()

        final_msg = f"Proceso completado. {success_count} de {total} PDFs generados exitosamente."
        self.logger.info(final_msg)
        if log_callback:
            log_callback(final_msg + "\n")
        return success_count, total

    def _fill_excel_template(self, template_path, output_path, data_row):
        """
        Lee el template, busca marcadores como {columna} o [columna] y los reemplaza.
        """
        wb = openpyxl.load_workbook(template_path)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        original_value = cell.value
                        new_value = original_value

                        for col_name, val in data_row.items():
                            placeholder_1 = f"{{{col_name}}}"
                            placeholder_2 = f"[{col_name}]"

                            if placeholder_1 in new_value:
                                new_value = new_value.replace(placeholder_1, str(val))
                            if placeholder_2 in new_value:
                                new_value = new_value.replace(placeholder_2, str(val))

                        if original_value != new_value:
                            cell.value = new_value

        wb.save(output_path)
        wb.close()

    def _convert_to_pdf(self, excel_app, excel_path, pdf_path):
        """
        Usa Excel para exportar el archivo xlsx a PDF.
        """
        wb = excel_app.Workbooks.Open(excel_path)
        try:
            wb.ExportAsFixedFormat(0, pdf_path)
        finally:
            wb.Close(False)