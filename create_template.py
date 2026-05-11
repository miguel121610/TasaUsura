import openpyxl
import os
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def create_example_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Formulario"

    # Estilos
    titulo_fill = PatternFill("solid", fgColor="1F4E78")
    seccion_fill = PatternFill("solid", fgColor="D9EAF7")
    label_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    fila = 1

    # Título principal
    ws["A1"] = "FORMULARIO CONSOLIDADO DE SOLICITUDES - RESTREPO"
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = titulo_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    fila = 3

    # =========================
    # 1. Solicitud
    # =========================
    ws[f"A{fila}"] = "1. Vista general de solicitudes con datos del cliente"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_1 = [
        ("ID Solicitud:", "{id_solicitud}"),
        ("Número cliente:", "{num_cliente}"),
        ("Nombres y apellidos:", "{nombres_apellidos}"),
        ("Cédula:", "{cedula}"),
        ("Teléfono celular:", "{telefono_celular}"),
        ("Ciudad:", "{ciudad}"),
        ("Monto solicitado:", "{monto_solicitado}"),
        ("Plazo meses:", "{plazo_meses}"),
        ("Cuota:", "{cuota}"),
        ("Ventas actual:", "{ventas_actual}"),
        ("Liquidez disponible:", "{liquidez_disponible}"),
        ("Estabilidad:", "{estabilidad}"),
        ("Habilidad empresarial:", "{habilidad_empresarial}"),
        ("Acuerdo socialización:", "{acuerdo_socializacion}"),
        ("Firma analista crédito:", "{firma_analista_credito}"),
        ("Fecha solicitud:", "{fecha_solicitud}"),
    ]

    for label, marcador in campos_1:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    fila += 1

    # =========================
    # 2. Ficha de visita
    # =========================
    ws[f"A{fila}"] = "2. Ficha de visita completa por solicitud"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_2 = [
        ("ID ficha:", "{id_ficha}"),
        ("Referencia:", "{ref}"),
        ("Cliente:", "{cliente}"),
        ("Cédula:", "{cedula}"),
        ("Actividad:", "{actividad}"),
        ("Nombre micro:", "{nombre_micro}"),
        ("Analista visita:", "{nombre_analista_visita}"),
        ("Tiempo antigüedad:", "{tiempo_antiguedad}"),
        ("Fecha desembolso:", "{fecha_desembolso}"),
        ("Destino crédito:", "{destino_credito}"),
        ("Lleva registros:", "{lleva_registros}"),
        ("Cómo se enteró del banco:", "{como_entero_banco}"),
        ("Fecha último pago:", "{fecha_ultimo_pago}"),
        ("Valor último pago:", "{valor_ultimo_pago}"),
        ("Lugar pago:", "{lugar_pago}"),
        ("Zona corresponde:", "{zona_corresponde}"),
        ("Observaciones:", "{observaciones}"),
    ]

    for label, marcador in campos_2:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    fila += 1

    # =========================
    # 3. Créditos activos
    # =========================
    ws[f"A{fila}"] = "3. Resumen de créditos activos por cliente"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_3 = [
        ("Nombres y apellidos:", "{nombres_apellidos}"),
        ("Cédula:", "{cedula}"),
        ("Número crédito:", "{numero_credito}"),
        ("Tipo:", "{tipo}"),
        ("Monto inicial:", "{monto_inicial}"),
        ("Saldo actual:", "{saldo_actual}"),
        ("Cuota mensual:", "{cuota_mensual}"),
        ("Meses:", "{meses}"),
        ("Fecha desembolso:", "{fecha_desembolso}"),
        ("Analista actual:", "{analista_actual}"),
        ("Analista cartera:", "{analista_cartera}"),
        ("Días mora:", "{dias_mora}"),
        ("Última fecha pago:", "{ultima_fecha_pago}"),
        ("Edad cliente:", "{edad_cliente}"),
        ("Atributo:", "{atributo}"),
    ]

    for label, marcador in campos_3:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    fila += 1

    # =========================
    # 4. Flujo financiero
    # =========================
    ws[f"A{fila}"] = "4. Análisis de flujo financiero (Anterior vs Actual)"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_4 = [
        ("Nombres y apellidos:", "{nombres_apellidos}"),
        ("Ventas anterior:", "{ventas_anterior}"),
        ("Ventas actual:", "{ventas_actual}"),
        ("Compras anterior:", "{compras_anterior}"),
        ("Compras actual:", "{compras_actual}"),
        ("Gastos generales anterior:", "{gastos_gen_ant}"),
        ("Gastos generales actual:", "{gastos_gen_act}"),
        ("Ingreso líquido anterior:", "{ing_liq_anterior}"),
        ("Ingreso líquido actual:", "{ing_liq_actual}"),
        ("Liquidez disponible anterior:", "{liq_disp_anterior}"),
        ("Liquidez disponible actual:", "{liq_disp_actual}"),
        ("Margen bruto anterior:", "{margen_bruto_ant}"),
        ("Margen bruto actual:", "{margen_bruto_act}"),
        ("Crecimiento ventas %:", "{crecimiento_ventas_pct}"),
        ("Crecimiento liquidez %:", "{crecimiento_liquidez_pct}"),
    ]

    for label, marcador in campos_4:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    fila += 1

    # =========================
    # 5. Clientes en mora
    # =========================
    ws[f"A{fila}"] = "5. Clientes en mora"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_5 = [
        ("Nombres y apellidos:", "{nombres_apellidos}"),
        ("Cédula:", "{cedula}"),
        ("Teléfono celular:", "{telefono_celular}"),
        ("Número crédito:", "{numero_credito}"),
        ("Tipo:", "{tipo}"),
        ("Saldo:", "{saldo}"),
        ("Días mora:", "{dias_mora}"),
        ("Última fecha pago:", "{ultima_fecha_pago}"),
        ("Analista cartera:", "{analista_cartera}"),
    ]

    for label, marcador in campos_5:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    fila += 1

    # =========================
    # 6. Detalle unificado
    # =========================
    ws[f"A{fila}"] = "6. Detalle completo unificado"
    ws.merge_cells(f"A{fila}:D{fila}")
    ws[f"A{fila}"].font = Font(bold=True)
    ws[f"A{fila}"].fill = seccion_fill
    fila += 1

    campos_6 = [
        ("ID solicitud:", "{id_solicitud}"),
        ("Número cliente:", "{num_cliente}"),
        ("Nombres y apellidos:", "{nombres_apellidos}"),
        ("Cédula:", "{cedula}"),
        ("Ciudad:", "{ciudad_ng}"),
        ("Actividad:", "{actividad}"),
        ("Nombre micro:", "{nombre_micro}"),
        ("Monto aprobado:", "{monto_aprobado}"),
        ("Plazo:", "{plazo}"),
        ("Cuota actual:", "{cuota_actual}"),
        ("Ventas actual:", "{ventas_actual}"),
        ("Liquidez disponible:", "{liq_disp_actual}"),
        ("Estabilidad:", "{estabilidad}"),
        ("Habilidad empresarial:", "{habilidad_empresarial}"),
        ("Días mora:", "{dias_mora}"),
        ("Analista actual:", "{analista_actual}"),
        ("Observaciones visita:", "{obs_visita}"),
        ("Observaciones solicitud:", "{obs_solicitud}"),
    ]

    for label, marcador in campos_6:
        ws[f"A{fila}"] = label
        ws[f"B{fila}"] = marcador
        ws[f"A{fila}"].fill = label_fill
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"A{fila}"].border = border
        ws[f"B{fila}"].border = border
        fila += 1

    # Ajustar ancho de columnas
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15

    template_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "plantilla_restrepo.xlsx"
    )
    wb.save(template_path)
    print(f"Plantilla Excel creada en: {template_path}")


if __name__ == "__main__":
    create_example_template()